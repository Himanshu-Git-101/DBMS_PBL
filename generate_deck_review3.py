"""
Hospital DBMS - Review 3 (Week 16 - 5 Marks) Presentation Generator
Generates widescreen 16:9 PowerPoint presentation for Final Review.
Concise, visual, high-contrast cards, and speaker notes on every slide.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_review3_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_DARK_BG    = RGBColor(11, 19, 41)      # #0B1329
    C_LIGHT_BG   = RGBColor(248, 250, 252)   # #F8FAFC
    C_CARD_BG    = RGBColor(255, 255, 255)   # #FFFFFF
    C_BORDER     = RGBColor(203, 213, 225)   # #CBD5E1
    
    C_TEAL       = RGBColor(13, 148, 136)    # #0D9488
    C_BLUE       = RGBColor(2, 132, 199)     # #0284C7
    C_PURPLE     = RGBColor(124, 58, 237)    # #7C3AED
    C_AMBER      = RGBColor(217, 119, 6)     # #D97706
    C_ROSE       = RGBColor(219, 39, 119)    # #DB2777
    C_GREEN      = RGBColor(5, 150, 105)     # #059669

    C_TEXT_DARK  = RGBColor(15, 23, 42)      # #0F172A
    C_TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B
    C_TEXT_LIGHT = RGBColor(255, 255, 255)   # #FFFFFF

    def set_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, badge_text, title_text, subtitle_text=None, is_dark=False):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(3.8), Inches(0.36))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(20, 184, 166) if is_dark else RGBColor(230, 255, 250)
        badge.line.color.rgb = C_TEAL
        tf = badge.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = badge_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = RGBColor(11, 19, 41) if is_dark else C_TEAL
        p.alignment = PP_ALIGN.CENTER

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.name = "Arial"
        p.font.color.rgb = C_TEXT_LIGHT if is_dark else C_TEXT_DARK

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(12)
            p2.font.color.rgb = RGBColor(148, 163, 184) if is_dark else C_TEXT_MUTED
            p2.space_before = Pt(2)

    def add_card(slide, x, y, w, h, border_top_color, title_text, items, is_dark=False):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(15, 23, 42) if is_dark else C_CARD_BG
        card.line.color.rgb = RGBColor(30, 41, 59) if is_dark else C_BORDER
        card.line.width = Pt(1.2)

        strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.12))
        strip.fill.solid()
        strip.fill.fore_color.rgb = border_top_color
        strip.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.16), Inches(w - 0.44), Inches(h - 0.28))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = border_top_color

        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(11)
            p.font.color.rgb = RGBColor(226, 232, 240) if is_dark else C_TEXT_DARK
            p.space_before = Pt(4)

        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE (Dark Executive Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, C_DARK_BG)

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(4.8), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_TEAL
    badge.line.fill.background()
    p = badge.text_frame.paragraphs[0]
    p.text = "DBMS PBL • REVIEW 3 (WEEK 16 — FINAL 5 MARKS)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER

    tb = s1.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Hospital Appointment &\nPatient Care Management System"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "Full Application Prototype, End-to-End CRUD Workflows, Error Handling & Final Project Demonstration"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_before = Pt(10)

    # 4 Highlight Badges
    badges = [
        ("FULL-STACK PROTOTYPE", "Zero-dependency Python application with Web & CLI", C_TEAL),
        ("END-TO-END CRUD", "Intake, booking, consultation, beds, billing & reports", C_BLUE),
        ("STRICT INTEGRITY", "Active triggers preventing overlaps & bed collisions", C_ROSE),
        ("VIVA & REPORT READY", "Full documentation, test logs, code & viva Q&A", C_AMBER)
    ]
    for i, (b_title, b_desc, b_color) in enumerate(badges):
        bx = 0.8 + i * 2.95
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(4.4), Inches(2.8), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(15, 23, 42)
        card.line.color.rgb = b_color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = b_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = b_color

        p2 = tf.add_paragraph()
        p2.text = b_desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(203, 213, 225)
        p2.space_before = Pt(8)

    s1.notes_slide.notes_text_frame.text = (
        "Review 3 Presenter Script:\n"
        "Respected evaluators, welcome to Review 3 — our final project demonstration for the Hospital Appointment "
        "and Patient Care Management System. Today we demonstrate our completed Python application connected directly "
        "to our 3NF relational database. We will walk through all clinical workflows, live CRUD operations, constraint "
        "enforcements, analytical reports, and our final viva defense."
    )

    # ==========================================
    # SLIDE 2: SYSTEM ARCHITECTURE
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, C_LIGHT_BG)
    add_header(s2, "Deliverable 1: Architecture", "Three-Tier Application Architecture", 
               "Modular separation between Presentation, Business Logic Controller, and Relational Database Engine")

    arch_cards = [
        ("Tier 1: Presentation Layer", [
            "Responsive Web Dashboard: Built with modern CSS (cards, glassmorphism, responsive grid).",
            "Real-Time User Interaction: Form validation, live patient search, and visual bed map.",
            "Interactive SQL Console: Integrated SQL sandbox for live query execution.",
            "Dual-Mode Interface: Also provides a lightweight terminal CLI mode for rapid evaluation."
        ], C_BLUE),
        ("Tier 2: Business Logic Engine", [
            "Zero-Dependency Python Controller: Built using Python standard library http.server.",
            "RESTful JSON API: Clean modular endpoints for Patients, Appointments, Encounters, Beds, and Bills.",
            "Transaction Management: Enforces ACID guarantees during multi-table writes.",
            "Error Interception: Translates low-level SQL constraint aborts into actionable UI alerts."
        ], C_PURPLE),
        ("Tier 3: Database Engine", [
            "Relational Database: Embedded SQLite / PostgreSQL with full 3NF normalization.",
            "Active Foreign Keys & Triggers: Schema-level collision detection and date checks.",
            "5 Precompiled Production Views: Abstracting complex multi-table joins.",
            "High-Performance Indexes: B-tree indexes on frequent join and search predicates."
        ], C_TEAL)
    ]
    for i, (title, items, color) in enumerate(arch_cards):
        add_card(s2, 0.8 + i * 3.95, 1.8, 3.8, 5.0, color, title, items)

    s2.notes_slide.notes_text_frame.text = (
        "Slide 2 Script: Architecture.\n"
        "Our system follows a standard three-tier architecture: "
        "The Presentation Layer provides an interactive, responsive web dashboard with live search and bed maps. "
        "The Business Logic Layer in Python coordinates API transactions and graceful error handling. "
        "The Relational Database Layer maintains 3NF tables, active triggers, and precompiled views."
    )

    # ==========================================
    # SLIDE 3: OUTPATIENT & CLINICAL ENCOUNTERS (CRUD)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, C_LIGHT_BG)
    add_header(s3, "Deliverable 2: Clinical CRUD", "Outpatient Registration & Clinical Workflows", 
               "Seamless end-to-end continuum: Patient Intake -> Appointment Booking -> Clinical Consultation")

    crud_cards = [
        ("1. Patient Registration & Search", [
            "Create (C): Master Patient Index intake with validation (phone uniqueness, DOB, blood group).",
            "Read (R): Real-time asynchronous search by name, phone number, or patient ID.",
            "Update (U): Patient demographics, emergency contacts, and residential addresses.",
            "Delete / Archive (D): Soft deactivation preserving historical medical continuity."
        ], C_BLUE),
        ("2. Overlap-Free Slot Booking", [
            "Duty Roster Integration: Validates doctor shifts and slot capacities.",
            "Overlap Collision Prevention: Trigger trg_prevent_appointment_overlap checks time intervals.",
            "Status Lifecycle: Transitioning from SCHEDULED -> CONFIRMED -> COMPLETED.",
            "Cancellation Handling: Re-opens slot immediately for incoming bookings."
        ], C_TEAL),
        ("3. Consultation & Diagnostics", [
            "1:1 Clinical Encounter: Links directly to appointment token with symptom recording.",
            "ICD-10 Diagnoses: Standard disease coding with severity levels (Mild to Critical).",
            "3NF Medication Prescriptions: Normalized line items (medicine, dosage, frequency, duration).",
            "Lab Test Orders: Orders diagnostic investigations directly to laboratory worklist."
        ], C_PURPLE)
    ]
    for i, (title, items, color) in enumerate(crud_cards):
        add_card(s3, 0.8 + i * 3.95, 1.8, 3.8, 5.0, color, title, items)

    s3.notes_slide.notes_text_frame.text = (
        "Slide 3 Script: Clinical CRUD.\n"
        "Here we demonstrate our core outpatient workflows. "
        "Patient intake establishes a unique Master Patient Index record. "
        "Appointment scheduling checks clinician rosters and enforces slot collision prevention. "
        "Doctors enter consultations with ICD-10 diagnoses, multi-item prescriptions, and lab test orders."
    )

    # ==========================================
    # SLIDE 4: INPATIENT ADMISSIONS & BILLING
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, C_LIGHT_BG)
    add_header(s4, "Deliverable 2: Inpatient & Billing", "Inpatient Bed Allocation & Billing Ledger", 
               "Ward census management with single-occupant bed locking and automated invoice settlement")

    ipd_cards = [
        ("Inpatient Bed Allocation", [
            "Interactive Bed Map: Color-coded visual grid (Green = Available, Red = Occupied).",
            "1 Active Patient Per Bed: Trigger trg_prevent_bed_double_booking aborts conflicting admissions.",
            "Attending Clinician Link: Connects patient, bed, ward, and admitting consultant.",
            "Discharge Workflow: Validates discharge_date >= admission_date and archives clinical summary."
        ], C_TEAL),
        ("Consolidated Hospital Billing", [
            "Inpatient & Outpatient Invoices: Consolidates consultation fees, room rent, and lab tests.",
            "Net Payable Formula: CHECK (net_payable = total_amount - discount_amount + tax_amount).",
            "Non-Negative Balances: All line items and totals guaranteed >= 0.",
            "Audit Trail: Invoice linked to consultation_id or admission_id."
        ], C_AMBER),
        ("Remittance & Payment Limits", [
            "Payment Methods: Multi-channel support (UPI, Credit Card, Debit Card, Cash, Insurance).",
            "Ceiling Constraint: CHECK (paid_amount <= net_payable) eliminates overpayments.",
            "Automated Trigger: trg_update_bill_on_payment increments receipts and sets PAID status.",
            "Itemized Receipt: Instant generation of printable financial vouchers."
        ], C_GREEN)
    ]
    for i, (title, items, color) in enumerate(ipd_cards):
        add_card(s4, 0.8 + i * 3.95, 1.8, 3.8, 5.0, color, title, items)

    s4.notes_slide.notes_text_frame.text = (
        "Slide 4 Script: Inpatient & Billing.\n"
        "In our inpatient module, beds are allocated with strict single-active-occupancy enforcement. "
        "The discharge process records clinical summaries and frees beds. "
        "Billing consolidates all fees into a single invoice where payments are capped by net payable."
    )

    # ==========================================
    # SLIDE 5: VALIDATION & ERROR HANDLING
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, C_LIGHT_BG)
    add_header(s5, "Deliverable 2: Validations", "Robust Validations & Defensive Error Handling", 
               "Dual-layer defense: Frontend form validation combined with unbreachable Database Triggers")

    val_cards = [
        ("Appointment Overlap Guard", [
            "Scenario: Booking Doctor 1 on 2026-10-10 at 09:10 when a 09:00-09:20 slot exists.",
            "Database Enforcement: Trigger raises ABORT with custom exception message.",
            "Application Response: Intercepts error and displays user-friendly collision alert.",
            "Result: 100% protection against doctor double-booking."
        ], C_ROSE),
        ("Bed Double-Booking Guard", [
            "Scenario: Admitting a second patient to an already occupied bed (e.g. ICU-BED-01).",
            "Database Enforcement: Trigger rejects insert with 'Bed currently has an active patient'.",
            "Application Response: UI highlights bed in red and blocks form submission.",
            "Result: Prevents physical bed occupancy conflicts."
        ], C_TEAL),
        ("Chronological & Financial Guards", [
            "Backwards Dates: CHECK constraint rejects discharge dates prior to admission intake.",
            "Payment Ceiling: Rejecting any payment that exceeds invoice net payable.",
            "Negative Values: Strict checks blocking negative consultation fees, costs, or quantities.",
            "Result: Preserves chronological and financial accounting integrity."
        ], C_AMBER)
    ]
    for i, (title, items, color) in enumerate(val_cards):
        add_card(s5, 0.8 + i * 3.95, 1.8, 3.8, 5.0, color, title, items)

    s5.notes_slide.notes_text_frame.text = (
        "Slide 5 Script: Error Handling.\n"
        "Our system features robust dual-layer validation. "
        "Frontend validations prevent obvious user mistakes, while database triggers and CHECK constraints "
        "provide an unbreachable barrier against slot overlaps, bed conflicts, date reversals, and payment overruns."
    )

    # ==========================================
    # SLIDE 6: COMPREHENSIVE REPORTING & 5 VIEWS
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, C_LIGHT_BG)
    add_header(s6, "Deliverable 2: Reports", "Analytical Reporting via 5 Production Views", 
               "Instant role-tailored reports delivering clinical insights and administrative intelligence")

    rep_cards = [
        ("Clinical Reporting", [
            "view_patient_history: Comprehensive longitudinal summary of visits, doctors seen, primary diagnoses, and count of medications & lab tests.",
            "view_pending_tests: Diagnostic laboratory worklist filtering tests awaiting collection or processing with turnaround hour alerts."
        ], C_BLUE),
        ("Operational Rostering", [
            "view_doctor_schedules: Real-time duty roster displaying shift hours, max capacity, currently booked appointments, and free slots.",
            "view_bed_occupancy: Live ward bed census computing operational beds, active inpatients, vacancy, and occupancy percentage."
        ], C_PURPLE),
        ("Financial Intelligence", [
            "view_outstanding_bills_revenue: Complete billing ledger tracking inpatient/outpatient invoices, gross revenue, collections, and pending balances.",
            "Interactive SQL Console: Live query runner enabling ad-hoc SQL execution and CSV export."
        ], C_GREEN)
    ]
    for i, (title, items, color) in enumerate(rep_cards):
        add_card(s6, 0.8 + i * 3.95, 1.8, 3.8, 5.0, color, title, items)

    s6.notes_slide.notes_text_frame.text = (
        "Slide 6 Script: Reporting.\n"
        "Our reporting dashboard provides one-click access to all 5 production database views. "
        "Clinicians access complete patient histories, lab technicians track pending tests, "
        "nursing staff monitor ward occupancy, and administrators inspect real-time revenue collection."
    )

    # ==========================================
    # SLIDE 7: TEST MATRIX & DEMONSTRATION VERIFICATION
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, C_LIGHT_BG)
    add_header(s7, "Deliverable 3: Test Verification", "CRUD & Integrity Test Execution Matrix", 
               "Systematic verification of all functional modules, triggers, and error recovery scenarios")

    test_cards = [
        ("Functional Test Cases (Pass)", [
            "TC-01: Patient registration with valid inputs -> Successfully created ID #11.",
            "TC-02: Search patient by phone '+91 91234 56780' -> Instant record match.",
            "TC-03: Booking appointment in free slot -> Token generated with status SCHEDULED.",
            "TC-04: Consultation entry with ICD-10 code -> Diagnosis & Rx items committed.",
            "TC-05: Inpatient bed allocation -> Bed status updated to Occupied."
        ], C_GREEN),
        ("Integrity Constraint Tests (Blocked)", [
            "TC-06: Overlapping appointment for Doctor 1 -> INTEGRITY ERROR raised & blocked.",
            "TC-07: Double-booking active Bed 1 -> BED ALLOCATION REJECTION blocked.",
            "TC-08: Discharge date prior to admission -> CHECK constraint failed & blocked.",
            "TC-09: Payment amount exceeding net bill -> CHECK constraint failed & blocked.",
            "TC-10: Negative consultation fee -> CHECK constraint failed & blocked."
        ], C_ROSE)
    ]
    for i, (title, items, color) in enumerate(test_cards):
        add_card(s7, 0.8 + i * 5.95, 1.8, 5.8, 5.0, color, title, items)

    s7.notes_slide.notes_text_frame.text = (
        "Slide 7 Script: Test Matrix.\n"
        "We executed a comprehensive 10-point test matrix. "
        "All positive functional CRUD tests passed flawlessly. "
        "Critically, all 5 negative test cases targeting slot overlaps, bed conflicts, backwards dates, "
        "and payment ceilings were actively blocked by our database triggers and constraints."
    )

    # ==========================================
    # SLIDE 8: CONCLUSION & EVALUATION READINESS
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, C_DARK_BG)
    add_header(s8, "Deliverable 3: Final Submission", "Project Summary & Evaluation Readiness", 
               "Complete delivery of all Review 3 University PBL guidelines (5/5 Marks Ready)", is_dark=True)

    summary_cards = [
        ("Delivered Project Assets", [
            "Working Prototype: Full Python web app and CLI mode (app.py).",
            "Production Relational DB: 3NF schema with triggers and sample data.",
            "Comprehensive Report: Full PBL dossier with ER, 3NF proofs, and SQL.",
            "Output Screens & Test Logs: Complete walkthrough and screenshots."
        ], C_TEAL),
        ("Viva-Voce Readiness", [
            "3NF Normalization: Mathematical proofs and functional dependency sets.",
            "ACID & Transactions: Rollback recovery and concurrency management.",
            "Indexing Strategy: B-Tree performance on foreign keys and lookups.",
            "Security & Views: Data encapsulation and role-based access."
        ], C_BLUE),
        ("Live Demonstration", [
            "1. Real-time patient intake and search.",
            "2. Overlap appointment collision prevention test.",
            "3. Live clinical consultation and ICD-10 entry.",
            "4. Bed allocation and single-patient lock demo.",
            "5. Instant 5-view reporting and SQL sandbox execution."
        ], C_AMBER)
    ]
    for i, (title, items, color) in enumerate(summary_cards):
        add_card(s8, 0.8 + i * 3.95, 1.8, 3.8, 5.0, color, title, items, is_dark=True)

    s8.notes_slide.notes_text_frame.text = (
        "Slide 8 Script: Conclusion.\n"
        "In conclusion, our Hospital Appointment and Patient Care Management System fulfills all Review 3 requirements. "
        "Our source code is clean, robust, and zero-dependency; our database enforces full relational integrity; "
        "and our documentation is comprehensive. We thank you for your mentorship and invite you to our live demonstration."
    )

    output_path = "Hospital_DBMS_Review3_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_review3_deck()
